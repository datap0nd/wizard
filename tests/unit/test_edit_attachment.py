"""Editing an attached workbook: Gemini's operations applied to a copy (openpyxl here, Excel on the work PC, checked with
fake Excel objects), the copy as a new file in the conversation to download, and the original never written."""
from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient
from tests.helpers import HEADERS, login, make_settings, wait_run

from wizard_api.app import create_app
from wizard_api.runs import ActiveRun, Recorder
from wizard_connectors.entitlements import IdentityDirectory
from wizard_connectors.paths import FIXTURES
from wizard_connectors.tools import ToolContext
from wizard_documents import edit as editing
from wizard_documents.edit import EditError, ExcelBook, OpenpyxlBook, apply, parse_value


def workbook(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    rows = [["Market", "Units", "Price", "Status"], ["EG", 120, 2.5, "Open"], ["SA", 340, 3, "Closed"],
            ["AE", 90, 4, "Open"], ["KW", 15, 5, "Closed"], ["QA", 200, 1.5, "Open"]]
    for row in rows:
        ws.append(row)
    ws["F1"] = "Note: EG is Egypt"
    wb.create_sheet("Setup")["A1"] = "FY start"
    wb.save(path)
    return path


def test_values_are_read_as_excel_would():
    assert parse_value("1234.5") == 1234.5 and parse_value("42") == 42 and parse_value("1,234") == 1234
    assert parse_value("=SUM(B2:B9)") == "=SUM(B2:B9)" and parse_value("'00123") == "00123" and parse_value("") is None
    assert str(parse_value("2026-10-23")) == "2026-10-23" and parse_value("TRUE") is True and parse_value("Open") == "Open"
    assert parse_value("1e3") == 1000.0 and parse_value("12 units") == "12 units"


def test_operations_change_a_copy_with_openpyxl(tmp_path):
    source = workbook(tmp_path / "sales.xlsx")
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    operations: list[dict[str, Any]] = [
        {"op": "set", "sheet": "Sales", "cell": "E1", "values": [["Revenue"], ["=B2*C2"], ["=B3*C3"], ["=B4*C4"], ["=B5*C5"], ["=B6*C6"]]},
        {"op": "delete_rows_where", "sheet": "Sales", "column": "Status", "condition": "equals", "value": "closed"},
        {"op": "sort", "sheet": "Sales", "range": "A1:E4", "column": "Units", "descending": True},
        {"op": "format", "sheet": "Sales", "range": "A1:E1", "bold": True, "fill": "#FFF2CC", "align": "center"},
        {"op": "format", "sheet": "Sales", "range": "C2:C4", "number_format": "#,##0.00"},
        {"op": "column_width", "sheet": "Sales", "columns": "A:B", "width": 18},
        {"op": "find_replace", "find": "EG", "replace": "Egypt", "whole_cell": True},
        {"op": "add_sheet", "name": "Summary", "after": "Sales"},
        {"op": "set", "sheet": "Summary", "cell": "A1", "values": [["Total units", "=SUM(Sales!B2:B4)"]]},
        {"op": "rename_sheet", "sheet": "Setup", "name": "Settings"},
        {"op": "freeze", "sheet": "Sales", "cell": "A2"},
        {"op": "autofilter", "sheet": "Sales", "range": "A1:E4"},
    ]
    result = editing.edit(source, operations, tmp_path / "out" / "sales (edited).xlsx", office="never")
    assert result["status"] == "ok" and result["method"] == "openpyxl"
    assert result["changes"][1] == "Sales: deleted 2 row(s) where D is equals 'closed'"
    assert result["changes"][6] == "replaced 'EG' with 'Egypt' in 1 cell(s)", "whole cell: the note in F1 keeps its EG"
    assert any("formulas that point at moved rows" in w for w in result["warnings"])
    assert result["sheets"] == ["Sales", "Summary", "Settings"]
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before, "the original is never written"
    wb = openpyxl.load_workbook(tmp_path / "out" / "sales (edited).xlsx")
    ws = wb["Sales"]
    assert [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=4, max_col=4)] == [
        ["Market", "Units", "Price", "Status"], ["QA", 200, 1.5, "Open"], ["Egypt", 120, 2.5, "Open"], ["AE", 90, 4, "Open"]]
    assert ws["E1"].value == "Revenue" and ws["F1"].value == "Note: EG is Egypt"
    assert ws["A1"].font.bold and ws["A1"].fill.start_color.rgb == "FFFFF2CC" and ws["A1"].alignment.horizontal == "center"
    assert ws["C2"].number_format == "#,##0.00" and ws.column_dimensions["A"].width == 18
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:E4"
    assert wb["Summary"]["B1"].value == "=SUM(Sales!B2:B4)" and wb.sheetnames == ["Sales", "Summary", "Settings"]


def test_mistakes_come_back_in_words(tmp_path):
    source = workbook(tmp_path / "sales.xlsx")
    book = OpenpyxlBook(source)
    for operations, words in (
        ([{"op": "set", "sheet": "Nope", "cell": "A1", "values": [["x"]]}], "no sheet 'Nope'; sheets: Sales, Setup"),
        ([{"op": "sort", "sheet": "Sales", "range": "A1:D6", "column": "Revenue"}], "no column 'Revenue'; headers: Market"),
        ([{"op": "set", "sheet": "Sales", "cell": "A0", "values": [["x"]]}], "'A0' is not a cell"),
        ([{"op": "format", "sheet": "Sales", "range": "A1", "fill": "yellow"}], "not a colour"),
        ([{"op": "add_sheet", "name": "Sales"}], "exists already"),
        ([{"op": "explode"}], "operation 1 (explode)"),
        ([], "no operations"),
    ):
        with pytest.raises(EditError, match=words.replace("(", r"\(").replace(")", r"\)")):
            apply(book, operations)
    with pytest.raises(EditError, match="need Excel"):
        OpenpyxlBook(tmp_path / "old.xls")


class Recording:
    """Records what is done to an Excel object; any attribute is another Recording."""

    def __init__(self, log: list[str], name: str, value: Any = None):
        self.__dict__.update(_log=log, _name=name, _value=value)

    def __getattr__(self, attribute: str) -> Any:
        return Recording(self._log, f"{self._name}.{attribute}")

    def __setattr__(self, attribute: str, value: Any) -> None:
        self._log.append(f"{self._name}.{attribute} = {value!r}")

    def __call__(self, *args: Any, **kwargs: Any) -> Any:
        self._log.append(f"{self._name}({', '.join([repr(a) for a in args] + [f'{k}={v!r}' for k, v in kwargs.items()])})")
        return Recording(self._log, f"{self._name}(…)")


class FakeSheet(Recording):
    def __init__(self, log: list[str], name: str, rows: list[list[Any]]):
        super().__init__(log, f"Worksheets({name!r})")
        self.__dict__.update(Name=name, rows=rows)

    @property
    def UsedRange(self) -> Any:  # noqa: N802 - Excel's name
        return type("Used", (), {"Row": 1, "Column": 1, "Rows": type("R", (), {"Count": len(self.rows)}),
                                 "Columns": type("C", (), {"Count": len(self.rows[0])})})()

    def Cells(self, row: int, column: int) -> tuple[int, int]:  # noqa: N802
        return row, column

    def Range(self, start: Any, end: Any = None) -> Any:  # noqa: N802
        sheet = self
        if not isinstance(start, tuple):
            return Recording(self._log, f"{self._name}.Range({start!r})")
        (top, left), (bottom, right) = start, end or start

        class Area(Recording):
            @property
            def Value(self) -> Any:  # noqa: N802
                return tuple(tuple(sheet.rows[r - 1][c - 1] for c in range(left, right + 1)) for r in range(top, bottom + 1))

        return Area(self._log, f"{self._name}.Range({top},{left}:{bottom},{right})")


def test_excel_edits_the_workbook_it_opened_read_only_and_saves_a_copy(tmp_path):
    log: list[str] = []
    rows = [["Market", "Units", "Status"], ["EG", 120, "Open"], ["SA", 340, "Closed"], ["AE", 90, "Open"], ["KW", 15, "Closed"]]
    sheet = FakeSheet(log, "Sales", rows)
    opened: list[dict[str, Any]] = []

    class Sheets:
        Count = 1

        def __call__(self, key: Any) -> FakeSheet:
            return sheet

    workbook_object = Recording(log, "wb")
    workbook_object.__dict__["Worksheets"] = Sheets()
    app = type("App", (), {"Workbooks": type("W", (), {"Open": staticmethod(lambda path, **options: opened.append(options)
                                                                          or workbook_object)})})()
    session = type("Session", (), {"app": lambda self, name: app, "error": lambda self, e: str(e)})()
    book = ExcelBook(session, tmp_path / "sales.xlsx")
    changes = apply(book, [
        {"op": "set", "sheet": "Sales", "cell": "D1", "values": [["Value"], ["=B2*2"]]},
        {"op": "delete_rows_where", "sheet": "Sales", "column": "Status", "condition": "equals", "value": "Closed"},
        {"op": "sort", "sheet": "Sales", "range": "A1:C3", "column": "B", "descending": True},
        {"op": "format", "sheet": "Sales", "range": "A1:D1", "bold": True, "fill": "#FF0000", "align": "center"},
        {"op": "column_width", "sheet": "Sales", "columns": "A:C"},
    ])
    book.save(tmp_path / "sales (edited).xlsx")
    book.close()
    assert opened == [{"UpdateLinks": 0, "ReadOnly": True, "IgnoreReadOnlyRecommended": True, "AddToMru": False}]
    assert changes[1] == "Sales: deleted 2 row(s) where C is equals 'Closed'"
    assert "Worksheets('Sales').Range(1,4:2,4).Formula = (('Value',), ('=B2*2',))" in log
    assert len([line for line in log if line.endswith(".Delete()")]) == 2, "one delete per block of matching rows"
    assert log.index(next(line for line in log if "Rows('5:5')" in line)) < log.index(next(line for line in log if "Rows('3:3')" in line)), \
        "rows are deleted from the bottom up so the row numbers stay valid"
    assert "Worksheets('Sales').Range(1,1:3,3).Sort(Key1=(1, 2), Order1=2, Header=1)" in log
    assert "Worksheets('Sales').Range(1,1:1,4).Interior.Color = 255" in log, "Excel colours are BGR"
    assert "Worksheets('Sales').Range(1,1:1,4).HorizontalAlignment = -4108" in log
    assert log.index("Worksheets('Sales').Columns('A:C')") == log.index("Worksheets('Sales').Columns(…).AutoFit()") - 1
    assert log[-2:] == [f"wb.SaveAs({str(tmp_path / 'sales (edited).xlsx')!r}, FileFormat=51)", "wb.Close(False)"]


def bound_file(client: TestClient, name: str, data: bytes) -> tuple[dict[str, Any], dict[str, Any]]:
    added = client.post("/api/v1/attachments", content=data,
                        headers={**HEADERS, "Content-Type": "application/octet-stream", "X-File-Name": name}).json()["attachment"]
    started = client.post("/api/v1/runs", json={"question": "Which market gave us the best return on marketing investment last quarter?",
                                                "attachment_ids": [added["id"]]}, headers=HEADERS).json()
    return started, wait_run(client, started["run_id"])


def test_the_edited_copy_is_a_new_file_to_download_and_read_back(tmp_path):
    source = workbook(tmp_path / "sales.xlsx")
    with TestClient(create_app(make_settings(tmp_path))) as client:
        login(client, "u-ceo")
        started, run = bound_file(client, "Gulf sales.xlsx", source.read_bytes())
        app = client.app
        assert run["status"] == "succeeded"
        ceo = IdentityDirectory.load(FIXTURES / "identities.json").get("u-ceo")
        assert ceo is not None
        active = ActiveRun("run_edit", ceo, started["conversation_id"], "ask", "edit it", None, "replay")
        recorder = Recorder(app.state.store, active, app.state.manager.attachments)
        context = ToolContext(identity=ceo, run_id="run_edit", recorder=recorder, services=app.state.manager.services)
        outcome = app.state.registry.execute("wizard_edit_attachment", {"file": "F1", "operations": [
            {"op": "delete_rows_where", "sheet": "Sales", "column": "Status", "condition": "equals", "value": "Closed"},
            {"op": "set", "sheet": "Sales", "cell": "E1", "values": [["Checked"], ["yes"], ["yes"], ["yes"]]}]}, context)
        assert outcome.ok, outcome.error_message
        assert outcome.data["file"] == "F2" and outcome.data["name"] == "Gulf sales (edited).xlsx"
        assert outcome.data["edited_with"].startswith("openpyxl") and outcome.data["readable"] is True
        assert active.files_made == [outcome.data["attachment_id"]]
        read = app.state.registry.execute("wizard_query_attachment", {"file": "F2", "sheet": "Sales", "limit": 10}, context)
        assert read.ok and read.data["total_rows"] == 3, "the edited copy is read back like any attached file"
        download = client.get(f"/api/v1/attachments/{outcome.data['attachment_id']}/file")
        assert download.status_code == 200 and "Gulf%20sales%20%28edited%29.xlsx" in download.headers["content-disposition"]
        edited = openpyxl.load_workbook(io.BytesIO(download.content))["Sales"]
        assert [r[0].value for r in edited.iter_rows(min_row=1, max_row=4)] == ["Market", "EG", "AE", "QA"]
        assert edited["E2"].value == "yes"
        original = app.state.store.conversation_attachments(started["conversation_id"])[0]
        assert hashlib.sha256(app.state.manager.attachments.original(original).read_bytes()).hexdigest() == \
            hashlib.sha256(source.read_bytes()).hexdigest()
        failed = app.state.registry.execute("wizard_edit_attachment", {"file": "F1", "operations": [
            {"op": "set", "sheet": "Missing", "cell": "A1", "values": [["x"]]}]}, context)
        assert not failed.ok and failed.error_code == "edit_failed" and "no sheet 'Missing'" in (failed.error_message or "")
        unknown = app.state.registry.execute("wizard_edit_attachment", {"file": "F9", "operations": [
            {"op": "clear", "sheet": "Sales", "range": "A1"}]}, context)
        assert not unknown.ok and unknown.error_code == "unknown_file"
        audit = [json.loads(r["detail"] or "{}") for r in app.state.store.db.execute(
            "SELECT detail FROM audit_log WHERE action='attachment:edit'").fetchall()]
        assert audit and audit[0]["operations"] == 2
        login(client, "u-cfo")
        assert client.get(f"/api/v1/attachments/{outcome.data['attachment_id']}/file").status_code == 404, "another user's file"


def test_gemini_is_told_to_edit_files_not_to_paste_code(registry):
    spec = registry.specs["wizard_edit_attachment"]
    assert "never answer with code" in spec.description and spec.category == "present"
    schema = json.dumps(registry.describe("wizard"))
    assert "delete_rows_where" in schema
