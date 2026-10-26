"""Edit a copy of a workbook the person attached: Gemini's list of operations, applied by Excel itself when Office is on
this PC (charts, pivots, formatting and NASCA-protected files survive, formulas follow moved cells) or by openpyxl
otherwise. The original is never written: Excel opens it read-only and saves the result under a new name.

Operations (all sheet and range names as the workbook shows them, ranges in A1 notation):
  set              sheet, cell, values (rows of cells: numbers as digits, formulas starting with =, text starting
                   with ' stays text)
  clear            sheet, range
  insert_rows      sheet, row, count          delete_rows     sheet, row, count
  insert_columns   sheet, column, count       delete_columns  sheet, column, count
  delete_rows_where sheet, column, condition (equals, not_equals, contains, not_contains, empty, not_empty, greater,
                   less), value, header_row
  sort             sheet, range, column (letter or header text), descending, header
  format           sheet, range, bold, italic, font_color, fill, number_format, align, wrap, border
  column_width     sheet, columns, width (or autofit)
  find_replace     sheet (or every sheet), find, replace, match_case, whole_cell
  add_sheet        name, after                rename_sheet    sheet, name       delete_sheet    sheet
  freeze           sheet, cell                autofilter      sheet, range"""
from __future__ import annotations

import contextlib
import copy
import datetime as dt
import re
import shutil
from pathlib import Path
from typing import Any, Protocol

from .common import ConversionError

EDITABLE = (".xlsx", ".xlsm", ".xls", ".xlsb")
OPENPYXL = (".xlsx", ".xlsm")
MAX_OPERATIONS = 60
MAX_CELLS = 20_000
CELL = re.compile(r"^\$?([A-Za-z]{1,3})\$?([0-9]{1,7})$")
COLUMN = re.compile(r"^\$?([A-Za-z]{1,3})$")
NUMBER = re.compile(r"^-?(\d+(\.\d*)?|\.\d+)([eE][-+]?\d+)?$")
ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
COLOR = re.compile(r"^#?[0-9A-Fa-f]{6}$")
FILE_FORMATS = {".xlsx": 51, ".xlsm": 52, ".xlsb": 50, ".xls": 56}  # Excel's xlOpenXMLWorkbook and friends
ALIGN = {"left": -4131, "center": -4108, "right": -4152}
CONDITIONS = ("equals", "not_equals", "contains", "not_contains", "empty", "not_empty", "greater", "less")


class EditError(ConversionError):
    pass


# --- Addresses and values ----------------------------------------------------------------------------------------------


def column_number(letters: str) -> int:
    number = 0
    for char in letters.upper():
        number = number * 26 + ord(char) - 64
    if not 1 <= number <= 16384:
        raise EditError(f"column {letters} is outside the sheet")
    return number


def column_letters(number: int) -> str:
    letters = ""
    while number:
        number, rest = divmod(number - 1, 26)
        letters = chr(65 + rest) + letters
    return letters


def cell(text: str) -> tuple[int, int]:
    match = CELL.match(str(text).strip())
    if not match or not 1 <= int(match.group(2)) <= 1_048_576:
        raise EditError(f"'{text}' is not a cell such as B3")
    return int(match.group(2)), column_number(match.group(1))


def cell_range(text: str) -> tuple[int, int, int, int]:
    """(first row, first column, last row, last column) of 'A1' or 'A1:C10'."""
    parts = str(text).strip().split(":")
    if len(parts) not in (1, 2):
        raise EditError(f"'{text}' is not a range such as A1:C10")
    top, left = cell(parts[0])
    bottom, right = cell(parts[-1])
    return min(top, bottom), min(left, right), max(top, bottom), max(left, right)


def columns_span(text: str) -> tuple[int, int]:
    parts = str(text).strip().split(":")
    numbers = []
    for part in parts:
        match = COLUMN.match(part)
        if not match:
            raise EditError(f"'{text}' is not a column or columns such as C or B:D")
        numbers.append(column_number(match.group(1)))
    return min(numbers), max(numbers)


def parse_value(text: Any) -> Any:
    """A cell as Gemini wrote it: '' or None empty, '=...' a formula, digits a number, ISO dates a date, a leading '
    keeps text as text."""
    if text is None or isinstance(text, bool | int | float):
        return text
    value = str(text)
    if value == "":
        return None
    if value.startswith("'"):
        return value[1:]
    if value.startswith("="):
        return value
    plain = value.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", value) else value
    if NUMBER.match(plain):
        number = float(plain)
        return int(number) if number.is_integer() and "." not in plain and "e" not in plain.lower() else number
    if ISO_DATE.match(value):
        with contextlib.suppress(ValueError):
            return dt.date.fromisoformat(value)
    if value.lower() in ("true", "false"):
        return value.lower() == "true"
    return value


def matches(value: Any, condition: str, wanted: str | None) -> bool:
    text = "" if value is None else str(value).strip()
    target = (wanted or "").strip()
    if condition == "empty":
        return text == ""
    if condition == "not_empty":
        return text != ""
    if condition in ("greater", "less"):
        try:
            number, limit = float(text.replace(",", "")), float(target.replace(",", ""))
        except ValueError:
            return False
        return number > limit if condition == "greater" else number < limit
    if condition in ("contains", "not_contains"):
        found = target.casefold() in text.casefold()
        return found if condition == "contains" else not found
    equal = text.casefold() == target.casefold()
    if not equal:
        with contextlib.suppress(ValueError):
            equal = float(text.replace(",", "")) == float(target.replace(",", ""))
    return equal if condition == "equals" else not equal


def color_hex(text: str) -> str:
    if not COLOR.match(text or ""):
        raise EditError(f"'{text}' is not a colour such as #FFF2CC")
    return text.lstrip("#").upper()


# --- Books ----------------------------------------------------------------------------------------------------------


class Book(Protocol):
    """What the operations need from a workbook; one implementation per engine."""

    method: str
    warnings: list[str]

    def sheets(self) -> list[str]: ...
    def size(self, sheet: str) -> tuple[int, int]: ...
    def read(self, sheet: str, top: int, left: int, bottom: int, right: int) -> list[list[Any]]: ...
    def write(self, sheet: str, top: int, left: int, rows: list[list[Any]]) -> None: ...
    def clear(self, sheet: str, top: int, left: int, bottom: int, right: int) -> None: ...
    def insert_rows(self, sheet: str, row: int, count: int) -> None: ...
    def delete_rows(self, sheet: str, row: int, count: int) -> None: ...
    def insert_columns(self, sheet: str, column: int, count: int) -> None: ...
    def delete_columns(self, sheet: str, column: int, count: int) -> None: ...
    def sort(self, sheet: str, area: tuple[int, int, int, int], key: int, descending: bool, header: bool) -> None: ...
    def style(self, sheet: str, area: tuple[int, int, int, int], spec: dict[str, Any]) -> None: ...
    def width(self, sheet: str, first: int, last: int, width: float | None) -> None: ...
    def replace(self, sheet: str, find: str, replace: str, match_case: bool, whole: bool) -> int: ...
    def add_sheet(self, name: str, after: str | None) -> None: ...
    def rename_sheet(self, sheet: str, name: str) -> None: ...
    def delete_sheet(self, sheet: str) -> None: ...
    def freeze(self, sheet: str, row: int, column: int) -> None: ...
    def autofilter(self, sheet: str, area: tuple[int, int, int, int]) -> None: ...
    def save(self, path: Path) -> None: ...
    def close(self) -> None: ...


def _sheet(book: Book, name: Any) -> str:
    names = book.sheets()
    wanted = str(name or "").strip()
    if not wanted and len(names) == 1:
        return names[0]
    found = next((n for n in names if n == wanted), None) or next((n for n in names if n.casefold() == wanted.casefold()), None)
    if found is None:
        raise EditError(f"no sheet '{wanted}'; sheets: {', '.join(names)}")
    return found


def _key_column(book: Book, sheet: str, area: tuple[int, int, int, int], column: str, header: bool) -> int:
    """A column named by its letter or by its header text, as a column number."""
    if COLUMN.match(column.strip()) and not (header and _header_index(book, sheet, area, column) is not None):
        number = column_number(column.strip().lstrip("$"))
        if not area[1] <= number <= area[3]:
            raise EditError(f"column {column} is outside {column_letters(area[1])}:{column_letters(area[3])}")
        return number
    found = _header_index(book, sheet, area, column)
    if found is None:
        headers = [str(v) for v in book.read(sheet, area[0], area[1], area[0], area[3])[0] if v not in (None, "")]
        raise EditError(f"no column '{column}'; headers: {', '.join(headers[:30])}")
    return found


def _header_index(book: Book, sheet: str, area: tuple[int, int, int, int], name: str) -> int | None:
    headers = book.read(sheet, area[0], area[1], area[0], area[3])[0]
    wanted = name.strip().casefold()
    for offset, value in enumerate(headers):
        if value is not None and str(value).strip().casefold() == wanted:
            return area[1] + offset
    return None


def _operate(book: Book, o: dict[str, Any]) -> str:  # noqa: C901 - one branch per operation, kept together
    op = o.get("op")
    if op == "add_sheet":
        name = str(o.get("name") or "").strip()
        if not name or len(name) > 31 or re.search(r"[\[\]:*?/\\]", name):
            raise EditError("a sheet name has 1-31 characters and none of [ ] : * ? / \\")
        if name.casefold() in (n.casefold() for n in book.sheets()):
            raise EditError(f"a sheet named '{name}' exists already")
        book.add_sheet(name, _sheet(book, o["after"]) if o.get("after") else None)
        return f"added sheet {name}"
    sheet = _sheet(book, o.get("sheet"))
    if op == "set":
        top, left = cell(o.get("cell") or "A1")
        rows = o.get("values")
        if rows is None and "value" in o:
            rows = [[o["value"]]]
        if not isinstance(rows, list) or not rows or not all(isinstance(r, list) for r in rows):
            raise EditError("values must be rows of cells, e.g. [[\"Total\", \"=SUM(B2:B9)\"]]")
        if sum(len(r) for r in rows) > MAX_CELLS:
            raise EditError(f"at most {MAX_CELLS:,} cells at a time")
        width = max(len(r) for r in rows)
        parsed = [[parse_value(v) for v in r] + [None] * (width - len(r)) for r in rows]
        book.write(sheet, top, left, parsed)
        end = f"{column_letters(left + width - 1)}{top + len(rows) - 1}"
        return f"{sheet}!{column_letters(left)}{top}" + (f":{end}" if (len(rows), width) != (1, 1) else "") + \
            f" set ({len(rows) * width} cell{'s' if len(rows) * width != 1 else ''})"
    if op == "clear":
        area = cell_range(o.get("range") or "")
        book.clear(sheet, *area)
        return f"{sheet}!{o['range']} cleared"
    if op in ("insert_rows", "delete_rows"):
        row, count = int(o.get("row") or 0), int(o.get("count") or 1)
        if not 1 <= row <= 1_048_576 or not 1 <= count <= 100_000:
            raise EditError("row must be a row number and count between 1 and 100,000")
        (book.insert_rows if op == "insert_rows" else book.delete_rows)(sheet, row, count)
        return f"{sheet}: {'inserted' if op == 'insert_rows' else 'deleted'} {count} row(s) at {row}"
    if op in ("insert_columns", "delete_columns"):
        column, count = column_number(str(o.get("column") or "").strip()), int(o.get("count") or 1)
        if not 1 <= count <= 1000:
            raise EditError("count must be between 1 and 1,000")
        (book.insert_columns if op == "insert_columns" else book.delete_columns)(sheet, column, count)
        return f"{sheet}: {'inserted' if op == 'insert_columns' else 'deleted'} {count} column(s) at {column_letters(column)}"
    if op == "delete_rows_where":
        condition = str(o.get("condition") or "equals")
        if condition not in CONDITIONS:
            raise EditError(f"condition must be one of {', '.join(CONDITIONS)}")
        header_row = int(o.get("header_row") or 1)
        rows, columns = book.size(sheet)
        area = (header_row, 1, max(rows, header_row), max(columns, 1))
        key = _key_column(book, sheet, area, str(o.get("column") or ""), True)
        values = book.read(sheet, header_row + 1, key, area[2], key) if area[2] > header_row else []
        hits = [header_row + 1 + i for i, row in enumerate(values) if matches(row[0], condition, o.get("value"))]
        for start, count in reversed(_runs(hits)):
            book.delete_rows(sheet, start, count)
        what = f"{condition.replace('_', ' ')} {o.get('value')!r}" if condition not in ("empty", "not_empty") else condition.replace("_", " ")
        return f"{sheet}: deleted {len(hits)} row(s) where {column_letters(key)} is {what}"
    if op == "sort":
        area = cell_range(o.get("range") or "")
        header = o.get("header", True) is not False
        key = _key_column(book, sheet, area, str(o.get("column") or column_letters(area[1])), header)
        book.sort(sheet, area, key, bool(o.get("descending")), header)
        return f"{sheet}!{o['range']} sorted by {column_letters(key)}{' descending' if o.get('descending') else ''}"
    if op == "format":
        area = cell_range(o.get("range") or "")
        spec = {k: o[k] for k in ("bold", "italic", "font_color", "fill", "number_format", "align", "wrap", "border")
                if o.get(k) is not None}
        if not spec:
            raise EditError("say what to format: bold, italic, font_color, fill, number_format, align, wrap or border")
        for colour in ("font_color", "fill"):
            if colour in spec:
                spec[colour] = color_hex(str(spec[colour]))
        if spec.get("align") not in (None, *ALIGN):
            raise EditError("align must be left, center or right")
        if spec.get("border") not in (None, "thin", "none"):
            raise EditError("border must be thin or none")
        book.style(sheet, area, spec)
        return f"{sheet}!{o['range']} formatted ({', '.join(spec)})"
    if op == "column_width":
        first, last = columns_span(str(o.get("columns") or ""))
        wanted = o.get("width")
        if wanted is not None and not 0 <= float(wanted) <= 255:
            raise EditError("width must be between 0 and 255")
        book.width(sheet, first, last, None if wanted is None else float(wanted))
        return f"{sheet}: columns {o['columns']} " + (f"width {wanted}" if wanted is not None else "fitted to their contents")
    if op == "rename_sheet":
        name = str(o.get("name") or "").strip()
        if not name or len(name) > 31 or re.search(r"[\[\]:*?/\\]", name):
            raise EditError("a sheet name has 1-31 characters and none of [ ] : * ? / \\")
        book.rename_sheet(sheet, name)
        return f"sheet {sheet} renamed {name}"
    if op == "delete_sheet":
        if len(book.sheets()) == 1:
            raise EditError("a workbook keeps at least one sheet")
        book.delete_sheet(sheet)
        return f"deleted sheet {sheet}"
    if op == "freeze":
        row, column = cell(o.get("cell") or "A2")
        book.freeze(sheet, row, column)
        return f"{sheet}: panes frozen at {column_letters(column)}{row}"
    if op == "autofilter":
        area = cell_range(o.get("range") or "")
        book.autofilter(sheet, area)
        return f"{sheet}!{o['range']}: filter buttons on"
    raise EditError(f"unknown operation '{op}'")


def _runs(rows: list[int]) -> list[tuple[int, int]]:
    """Consecutive row numbers as (first, count) blocks: deleting blocks bottom-up keeps the numbers valid."""
    blocks: list[tuple[int, int]] = []
    for row in rows:
        if blocks and blocks[-1][0] + blocks[-1][1] == row:
            blocks[-1] = (blocks[-1][0], blocks[-1][1] + 1)
        else:
            blocks.append((row, 1))
    return blocks


def apply(book: Book, operations: list[dict[str, Any]]) -> list[str]:
    """Run the operations in order; each returns what it changed, in words. find_replace covers one sheet or every
    sheet."""
    if not operations:
        raise EditError("no operations")
    if len(operations) > MAX_OPERATIONS:
        raise EditError(f"at most {MAX_OPERATIONS} operations at a time")
    changes = []
    for index, operation in enumerate(operations, 1):
        try:
            if operation.get("op") == "find_replace":
                find = str(operation.get("find") or "")
                if not find:
                    raise EditError("find is empty")
                sheets = [_sheet(book, operation["sheet"])] if operation.get("sheet") else book.sheets()
                count = sum(book.replace(s, find, str(operation.get("replace") or ""), bool(operation.get("match_case")),
                                         bool(operation.get("whole_cell"))) for s in sheets)
                changes.append(f"replaced '{find}' with '{operation.get('replace') or ''}' in {count} cell(s)")
            else:
                changes.append(_operate(book, operation))
        except EditError as error:
            raise EditError(f"operation {index} ({operation.get('op')}): {error}") from None
        except Exception as error:  # noqa: BLE001 - openpyxl or Excel refused it: say which operation, in words
            raise EditError(f"operation {index} ({operation.get('op')}) could not be done: "
                            f"{str(error).strip()[:300] or type(error).__name__}") from None
    return changes


# --- openpyxl ---------------------------------------------------------------------------------------------------------


class OpenpyxlBook:
    """A workbook edited with openpyxl (no Office). Conditions read the values Excel last saved for formula cells."""

    method = "openpyxl"

    def __init__(self, path: Path):
        try:
            import openpyxl
        except ImportError:
            raise EditError("openpyxl is not installed, and Office is not available to edit the file") from None
        if path.suffix.lower() not in OPENPYXL:
            raise EditError(f"{path.suffix} files need Excel to be edited; this PC cannot start it")
        try:
            self.wb = openpyxl.load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
            self.values = openpyxl.load_workbook(path, data_only=True, read_only=False)
        except Exception as error:  # noqa: BLE001 - zip, XML or encryption errors, in words
            raise EditError(f"the file could not be opened without Excel ({type(error).__name__}); it may be protected"
                            ) from None
        self.warnings: list[str] = []
        if any(getattr(ws, "_charts", None) or getattr(ws, "_images", None) for ws in self.wb.worksheets):
            self.warnings.append("Edited without Excel: charts and pictures in this workbook may not be kept.")

    def sheets(self) -> list[str]:
        return list(self.wb.sheetnames)

    def size(self, sheet: str) -> tuple[int, int]:
        ws = self.wb[sheet]
        return ws.max_row, ws.max_column

    def read(self, sheet: str, top: int, left: int, bottom: int, right: int) -> list[list[Any]]:
        values = self.values[sheet] if sheet in self.values.sheetnames else self.wb[sheet]
        own = self.wb[sheet]
        rows = []
        for r in range(top, bottom + 1):
            row = []
            for c in range(left, right + 1):
                value = values.cell(r, c).value
                if value is None:  # a cell written in this edit has no saved value yet
                    value = own.cell(r, c).value
                row.append(value)
            rows.append(row)
        return rows

    def write(self, sheet: str, top: int, left: int, rows: list[list[Any]]) -> None:
        ws, shadow = self.wb[sheet], self.values[sheet]
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                ws.cell(top + i, left + j).value = value
                shadow.cell(top + i, left + j).value = None if isinstance(value, str) and value.startswith("=") else value

    def clear(self, sheet: str, top: int, left: int, bottom: int, right: int) -> None:
        for book in (self.wb, self.values):
            ws = book[sheet]
            for r in range(top, bottom + 1):
                for c in range(left, right + 1):
                    ws.cell(r, c).value = None

    def _moved(self) -> None:
        note = "Edited without Excel: formulas that point at moved rows or columns were not updated."
        if note not in self.warnings:
            self.warnings.append(note)

    def insert_rows(self, sheet: str, row: int, count: int) -> None:
        for book in (self.wb, self.values):
            book[sheet].insert_rows(row, count)
        self._moved()

    def delete_rows(self, sheet: str, row: int, count: int) -> None:
        for book in (self.wb, self.values):
            book[sheet].delete_rows(row, count)
        self._moved()

    def insert_columns(self, sheet: str, column: int, count: int) -> None:
        for book in (self.wb, self.values):
            book[sheet].insert_cols(column, count)
        self._moved()

    def delete_columns(self, sheet: str, column: int, count: int) -> None:
        for book in (self.wb, self.values):
            book[sheet].delete_cols(column, count)
        self._moved()

    def sort(self, sheet: str, area: tuple[int, int, int, int], key: int, descending: bool, header: bool) -> None:
        top, left, bottom, right = area
        first = top + 1 if header else top
        if bottom < first:
            return
        ws, shadow = self.wb[sheet], self.values[sheet]
        rows = [[ws.cell(r, c).value for c in range(left, right + 1)] for r in range(first, bottom + 1)]
        saved = [[shadow.cell(r, c).value for c in range(left, right + 1)] for r in range(first, bottom + 1)]
        styles = [[ws.cell(r, c)._style for c in range(left, right + 1)] for r in range(first, bottom + 1)]

        def sort_key(index: int) -> tuple[int, Any]:
            value = saved[index][key - left] if saved[index][key - left] is not None else rows[index][key - left]
            if value is None or value == "":
                return (2, "")
            if isinstance(value, int | float) and not isinstance(value, bool):
                return (0, value)
            return (1, str(value).casefold())

        blanks = [i for i in range(len(rows)) if sort_key(i)[0] == 2]
        filled = sorted((i for i in range(len(rows)) if sort_key(i)[0] != 2), key=sort_key, reverse=descending)
        for target, source in enumerate(filled + blanks):
            for j in range(right - left + 1):
                ws.cell(first + target, left + j).value = rows[source][j]
                ws.cell(first + target, left + j)._style = styles[source][j]
                shadow.cell(first + target, left + j).value = saved[source][j]
        if any(isinstance(v, str) and v.startswith("=") for row in rows for v in row):
            self.warnings.append("Edited without Excel: sorted formulas keep their original references.")

    def style(self, sheet: str, area: tuple[int, int, int, int], spec: dict[str, Any]) -> None:
        from openpyxl.styles import Border, PatternFill, Side
        ws = self.wb[sheet]
        top, left, bottom, right = area
        if (bottom - top + 1) * (right - left + 1) > 200_000:
            raise EditError("format at most 200,000 cells at a time")
        thin = Side(style="thin")
        for r in range(top, bottom + 1):
            for c in range(left, right + 1):
                target = ws.cell(r, c)
                if any(k in spec for k in ("bold", "italic", "font_color")):
                    font = copy.copy(target.font)
                    font.bold = spec.get("bold", font.bold)
                    font.italic = spec.get("italic", font.italic)
                    if "font_color" in spec:
                        font.color = "FF" + spec["font_color"]
                    target.font = font
                if "fill" in spec:
                    target.fill = PatternFill("solid", start_color="FF" + spec["fill"], end_color="FF" + spec["fill"])
                if "number_format" in spec:
                    target.number_format = str(spec["number_format"])
                if "align" in spec or "wrap" in spec:
                    alignment = copy.copy(target.alignment)
                    alignment.horizontal = spec.get("align", alignment.horizontal)
                    alignment.wrap_text = spec.get("wrap", alignment.wrap_text)
                    target.alignment = alignment
                if "border" in spec:
                    target.border = Border(left=thin, right=thin, top=thin, bottom=thin) if spec["border"] == "thin" else Border()

    def width(self, sheet: str, first: int, last: int, width: float | None) -> None:
        ws = self.wb[sheet]
        for column in range(first, last + 1):
            if width is None:
                longest = max((len(str(v)) for (v,) in ws.iter_rows(min_col=column, max_col=column, values_only=True)
                               if v is not None), default=8)
                ws.column_dimensions[column_letters(column)].width = min(80, longest * 1.15 + 2)
            else:
                ws.column_dimensions[column_letters(column)].width = width

    def replace(self, sheet: str, find: str, replace: str, match_case: bool, whole: bool) -> int:
        ws, shadow = self.wb[sheet], self.values[sheet]
        count = 0
        flags = 0 if match_case else re.IGNORECASE
        for row in ws.iter_rows():
            for target in row:
                value = target.value
                if not isinstance(value, str) or value.startswith("="):
                    continue
                if whole:
                    if (value == find) if match_case else (value.casefold() == find.casefold()):
                        new: Any = parse_value(replace) if replace else None
                    else:
                        continue
                else:
                    new = re.sub(re.escape(find), lambda _: replace, value, flags=flags)
                    if new == value:
                        continue
                target.value = new
                shadow.cell(target.row, target.column).value = new
                count += 1
        return count

    def add_sheet(self, name: str, after: str | None) -> None:
        index = self.wb.sheetnames.index(after) + 1 if after else len(self.wb.sheetnames)
        self.wb.create_sheet(name, index)
        self.values.create_sheet(name, index)

    def rename_sheet(self, sheet: str, name: str) -> None:
        self.wb[sheet].title = name
        self.values[sheet].title = name

    def delete_sheet(self, sheet: str) -> None:
        self.wb.remove(self.wb[sheet])
        self.values.remove(self.values[sheet])

    def freeze(self, sheet: str, row: int, column: int) -> None:
        self.wb[sheet].freeze_panes = f"{column_letters(column)}{row}"

    def autofilter(self, sheet: str, area: tuple[int, int, int, int]) -> None:
        top, left, bottom, right = area
        self.wb[sheet].auto_filter.ref = f"{column_letters(left)}{top}:{column_letters(right)}{bottom}"

    def save(self, path: Path) -> None:
        self.wb.save(path)

    def close(self) -> None:
        with contextlib.suppress(Exception):
            self.wb.close()
            self.values.close()


# --- Excel (Office COM) ------------------------------------------------------------------------------------------------


class ExcelBook:
    """A workbook edited by Excel itself: opened read-only (the original is never written), saved under a new name."""

    method = "excel"

    def __init__(self, session: Any, path: Path):
        self.session, self.path = session, path
        self.app = session.app("excel")
        try:
            self.wb = self.app.Workbooks.Open(str(path), UpdateLinks=0, ReadOnly=True, IgnoreReadOnlyRecommended=True,
                                              AddToMru=False)
        except Exception as error:  # noqa: BLE001 - COM errors, in words
            raise EditError(f"Excel could not open it: {session.error(error)}") from None
        self.warnings: list[str] = []

    def _ws(self, sheet: str) -> Any:
        return self.wb.Worksheets(sheet)

    def _range(self, sheet: str, top: int, left: int, bottom: int, right: int) -> Any:
        ws = self._ws(sheet)
        return ws.Range(ws.Cells(top, left), ws.Cells(bottom, right))

    def sheets(self) -> list[str]:
        count = int(self.wb.Worksheets.Count)
        return [str(self.wb.Worksheets(i).Name) for i in range(1, count + 1)]

    def size(self, sheet: str) -> tuple[int, int]:
        used = self._ws(sheet).UsedRange
        return int(used.Row) + int(used.Rows.Count) - 1, int(used.Column) + int(used.Columns.Count) - 1

    def read(self, sheet: str, top: int, left: int, bottom: int, right: int) -> list[list[Any]]:
        value = self._range(sheet, top, left, bottom, right).Value
        if not isinstance(value, tuple):
            return [[value]]
        return [list(row) for row in value]

    def write(self, sheet: str, top: int, left: int, rows: list[list[Any]]) -> None:
        converted = tuple(tuple(v.isoformat() if isinstance(v, dt.date) else v for v in row) for row in rows)
        target = self._range(sheet, top, left, top + len(rows) - 1, left + len(rows[0]) - 1)
        target.Formula = converted[0][0] if len(converted) == 1 and len(converted[0]) == 1 else converted

    def clear(self, sheet: str, top: int, left: int, bottom: int, right: int) -> None:
        self._range(sheet, top, left, bottom, right).ClearContents()

    def insert_rows(self, sheet: str, row: int, count: int) -> None:
        self._ws(sheet).Rows(f"{row}:{row + count - 1}").Insert()

    def delete_rows(self, sheet: str, row: int, count: int) -> None:
        self._ws(sheet).Rows(f"{row}:{row + count - 1}").Delete()

    def insert_columns(self, sheet: str, column: int, count: int) -> None:
        self._ws(sheet).Columns(f"{column_letters(column)}:{column_letters(column + count - 1)}").Insert()

    def delete_columns(self, sheet: str, column: int, count: int) -> None:
        self._ws(sheet).Columns(f"{column_letters(column)}:{column_letters(column + count - 1)}").Delete()

    def sort(self, sheet: str, area: tuple[int, int, int, int], key: int, descending: bool, header: bool) -> None:
        top, left, bottom, right = area
        ws = self._ws(sheet)
        self._range(sheet, *area).Sort(Key1=ws.Cells(top, key), Order1=2 if descending else 1, Header=1 if header else 2)

    def style(self, sheet: str, area: tuple[int, int, int, int], spec: dict[str, Any]) -> None:
        target = self._range(sheet, *area)

        def bgr(hex_color: str) -> int:
            red, green, blue = int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
            return red + green * 256 + blue * 65536

        if "bold" in spec:
            target.Font.Bold = bool(spec["bold"])
        if "italic" in spec:
            target.Font.Italic = bool(spec["italic"])
        if "font_color" in spec:
            target.Font.Color = bgr(spec["font_color"])
        if "fill" in spec:
            target.Interior.Color = bgr(spec["fill"])
        if "number_format" in spec:
            target.NumberFormat = str(spec["number_format"])
        if "align" in spec:
            target.HorizontalAlignment = ALIGN[spec["align"]]
        if "wrap" in spec:
            target.WrapText = bool(spec["wrap"])
        if "border" in spec:
            target.Borders.LineStyle = 1 if spec["border"] == "thin" else -4142  # xlContinuous / xlNone

    def width(self, sheet: str, first: int, last: int, width: float | None) -> None:
        columns = self._ws(sheet).Columns(f"{column_letters(first)}:{column_letters(last)}")
        if width is None:
            columns.AutoFit()
        else:
            columns.ColumnWidth = width

    def replace(self, sheet: str, find: str, replace: str, match_case: bool, whole: bool) -> int:
        ws = self._ws(sheet)
        used = ws.UsedRange
        before = self.read(sheet, int(used.Row), int(used.Column), *self.size(sheet))
        used.Replace(What=find, Replacement=replace, LookAt=1 if whole else 2, MatchCase=match_case)
        after = self.read(sheet, int(used.Row), int(used.Column), *self.size(sheet))
        return sum(1 for old_row, new_row in zip(before, after, strict=False) for old, new in zip(old_row, new_row, strict=False)
                   if old != new)

    def add_sheet(self, name: str, after: str | None) -> None:
        anchor = self._ws(after) if after else self.wb.Worksheets(int(self.wb.Worksheets.Count))
        added = self.wb.Worksheets.Add(After=anchor)
        added.Name = name

    def rename_sheet(self, sheet: str, name: str) -> None:
        self._ws(sheet).Name = name

    def delete_sheet(self, sheet: str) -> None:
        self._ws(sheet).Delete()

    def freeze(self, sheet: str, row: int, column: int) -> None:
        try:
            self._ws(sheet).Activate()
            window = self.app.ActiveWindow
            window.FreezePanes = False
            window.SplitColumn, window.SplitRow = column - 1, row - 1
            window.FreezePanes = True
        except Exception:  # noqa: BLE001 - no window for a hidden workbook
            self.warnings.append(f"Panes on {sheet} could not be frozen here; freeze them in Excel.")

    def autofilter(self, sheet: str, area: tuple[int, int, int, int]) -> None:
        ws = self._ws(sheet)
        if not bool(getattr(ws, "AutoFilterMode", False)):
            self._range(sheet, *area).AutoFilter()

    def save(self, path: Path) -> None:
        try:
            self.wb.SaveAs(str(path), FileFormat=FILE_FORMATS.get(path.suffix.lower(), 51))
        except Exception as error:  # noqa: BLE001 - COM errors, in words
            raise EditError(f"Excel could not save the edited copy: {self.session.error(error)}") from None

    def save_copy(self, path: Path) -> None:
        try:
            self.wb.SaveCopyAs(str(path))
        except Exception as error:  # noqa: BLE001 - COM errors, in words
            raise EditError(f"Excel could not save the copy: {self.session.error(error)}") from None

    def close(self) -> None:
        with contextlib.suppress(Exception):
            self.wb.Close(False)


def open_book(path: Path, session: Any = None) -> Book:
    """Excel through `session` when Office runs here, else openpyxl."""
    if path.suffix.lower() not in EDITABLE:
        raise EditError(f"{path.suffix or 'this file type'} cannot be edited; Wizard edits Excel workbooks")
    if session is not None:
        return ExcelBook(session, path)
    return OpenpyxlBook(path)


def edit(source: Path, operations: list[dict[str, Any]], output: Path, office: str = "auto") -> dict[str, Any]:
    """Apply the operations to a copy of `source` saved as `output`; returns what changed."""
    from .office import OfficeSession, OfficeUnavailable
    session = None
    if office != "never":
        try:
            session = OfficeSession()
        except OfficeUnavailable as error:
            if office == "always":
                raise EditError(str(error)) from None
    try:
        book = open_book(source, session)
        try:
            changes = apply(book, operations)
            sheets = book.sheets()
            output.parent.mkdir(parents=True, exist_ok=True)
            book.save(output)
        finally:
            book.close()
        return {"status": "ok", "method": book.method, "changes": changes, "warnings": book.warnings, "sheets": sheets}
    finally:
        if session is not None:
            session.close()


def save_copy(source: Path, target: Path, office: str = "auto") -> dict[str, Any]:
    """Put a copy of a workbook Wizard made at `target` (the person's Downloads folder). Excel writes it when Office runs
    here: a NASCA-protected workbook opens only where Office itself saved it, and a browser download of one can fail
    ("Couldn't finish download"). Without Office, or when Excel cannot, the bytes are copied. Never overwrites."""
    from .office import OfficeSession, OfficeUnavailable
    if target.exists():
        raise EditError(f"{target.name} already exists")
    target.parent.mkdir(parents=True, exist_ok=True)
    problem = None
    if office != "never":
        try:
            with OfficeSession() as session:
                book = ExcelBook(session, source)
                try:
                    book.save_copy(target)
                finally:
                    book.close()
            if target.is_file():
                return {"status": "ok", "method": "excel", "warnings": []}
            problem = "Excel did not write the copy"
        except OfficeUnavailable:
            pass
        except EditError as error:
            problem = str(error)
    shutil.copyfile(source, target)
    return {"status": "ok", "method": "copy",
            "warnings": [f"{problem}; Wizard copied the file as it is, so a protected workbook may not open."] if problem else []}
